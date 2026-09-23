"""
Loads SOP documents (PDF, DOCX, Markdown) from a directory, chunks them, and
returns LangChain-style Documents ready to be embedded into ChromaDB.
"""
from __future__ import annotations

from dataclasses import dataclass, field
from pathlib import Path
from typing import List

from langchain.text_splitter import RecursiveCharacterTextSplitter
from pypdf import PdfReader
from docx import Document as DocxDocument

from logging_setup import logger

SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".md", ".txt"}


@dataclass
class SopDocument:
    """A chunk of a source SOP file, ready for embedding."""

    content: str
    source: str
    doc_type: str
    metadata: dict = field(default_factory=dict)


def _read_pdf(path: Path) -> str:
    reader = PdfReader(str(path))
    return "\n".join(page.extract_text() or "" for page in reader.pages)


def _read_docx(path: Path) -> str:
    doc = DocxDocument(str(path))
    return "\n".join(p.text for p in doc.paragraphs if p.text.strip())


def _read_text(path: Path) -> str:
    return path.read_text(encoding="utf-8", errors="ignore")


_READERS = {
    ".pdf": _read_pdf,
    ".docx": _read_docx,
    ".md": _read_text,
    ".txt": _read_text,
}


def load_sop_directory(directory: str | Path, chunk_size: int = 800, chunk_overlap: int = 120) -> List[SopDocument]:
    """Walk a directory of SOP files and return chunked SopDocuments."""
    directory = Path(directory)
    if not directory.exists():
        logger.warning(f"SOP directory does not exist: {directory}")
        return []

    splitter = RecursiveCharacterTextSplitter(
        chunk_size=chunk_size,
        chunk_overlap=chunk_overlap,
        separators=["\n\n", "\n", ". ", " ", ""],
    )

    documents: List[SopDocument] = []
    files = [p for p in directory.rglob("*") if p.suffix.lower() in SUPPORTED_EXTENSIONS]
    logger.info(f"Found {len(files)} SOP files in {directory}")

    for path in files:
        reader = _READERS.get(path.suffix.lower())
        if reader is None:
            continue
        try:
            raw_text = reader(path)
        except Exception as exc:  # noqa: BLE001 - log and continue on a bad file
            logger.error(f"Failed to read {path}: {exc}")
            continue

        if not raw_text.strip():
            logger.warning(f"No extractable text in {path}")
            continue

        chunks = splitter.split_text(raw_text)
        for idx, chunk in enumerate(chunks):
            documents.append(
                SopDocument(
                    content=chunk,
                    source=path.name,
                    doc_type=path.suffix.lower().lstrip("."),
                    metadata={
                        "source_path": str(path),
                        "chunk_index": idx,
                        "total_chunks": len(chunks),
                    },
                )
            )

    logger.info(f"Produced {len(documents)} SOP chunks from {len(files)} files")
    return documents


if __name__ == "__main__":
    import sys

    target = sys.argv[1] if len(sys.argv) > 1 else "sample_data/sops"
    docs = load_sop_directory(target)
    for d in docs[:3]:
        print(f"--- {d.source} (chunk {d.metadata['chunk_index']}) ---")
        print(d.content[:200], "...\n")
